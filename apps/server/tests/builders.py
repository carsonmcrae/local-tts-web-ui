"""Purpose: Build realistic (and deliberately malicious) test documents in memory.

Scope: Minimal hand-written PDFs (text pages, outline, image-only pages), DOCX via python-docx,
EPUBs via zipfile, and hostile archives (zip bombs, XML entity expansion, path traversal names).
Generating fixtures in code keeps binary blobs out of the repository and documents exactly what
each case contains.
Does NOT belong: Assertions (tests) or app code.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx


def _pdf_string(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(
    pages: list[list[str]],
    outline: list[tuple[str, int]] | None = None,
    image_only: bool = False,
    title: str | None = None,
) -> bytes:
    """A valid PDF with one text line per list item (Helvetica 11pt), optional outline and image pages."""
    numbered: dict[int, bytes] = {}
    page_ids: list[int] = []
    font_id = 3
    first_page_id = 4
    for index, lines in enumerate(pages):
        page_id = first_page_id + index * 2
        content_id = page_id + 1
        page_ids.append(page_id)
        if image_only:
            stream = b"q 400 0 0 600 100 100 cm 0.5 g 0 0 1 1 re f Q"
        else:
            ops = ["BT", "/F1 11 Tf", "14 TL", "72 760 Td", *[f"({_pdf_string(line)}) Tj T*" for line in lines], "ET"]
            stream = "\n".join(ops).encode("latin-1")
        numbered[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id
        )
        numbered[content_id] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    numbered[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    numbered[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    catalog = "<< /Type /Catalog /Pages 2 0 R"
    if outline:
        outlines_id = first_page_id + len(pages) * 2
        item_ids = [outlines_id + 1 + i for i in range(len(outline))]
        numbered[outlines_id] = (
            f"<< /Type /Outlines /First {item_ids[0]} 0 R /Last {item_ids[-1]} 0 R /Count {len(item_ids)} >>".encode()
        )
        for i, (entry_title, page_index) in enumerate(outline):
            links = ""
            if i > 0:
                links += f" /Prev {item_ids[i - 1]} 0 R"
            if i < len(item_ids) - 1:
                links += f" /Next {item_ids[i + 1]} 0 R"
            numbered[item_ids[i]] = (
                f"<< /Title ({_pdf_string(entry_title)}) /Parent {outlines_id} 0 R /Dest [{page_ids[page_index]} 0 R /Fit]{links} >>"
            ).encode()
        catalog += f" /Outlines {outlines_id} 0 R"
    numbered[1] = (catalog + " >>").encode()

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(numbered):
        offsets[number] = out.tell()
        out.write(f"{number} 0 obj\n".encode() + numbered[number] + b"\nendobj\n")
    xref = out.tell()
    size = max(numbered) + 1
    out.write(f"xref\n0 {size}\n0000000000 65535 f \n".encode())
    for number in range(1, size):
        out.write(f"{offsets.get(number, 0):010d} 00000 n \n".encode())
    info = f" /Info << /Title ({_pdf_string(title)}) >>" if title else ""
    out.write(f"trailer\n<< /Size {size} /Root 1 0 R{info} >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(
    path: Path, blocks: list[tuple[str, str]], title: str | None = None, table: list[list[str]] | None = None
) -> Path:
    """``blocks`` are ``(style, text)`` pairs, e.g. ``("Heading 1", "Intro")`` or ``("Normal", "…")``."""
    document = docx.Document()
    if title:
        document.core_properties.title = title
    for style, text in blocks:
        if style.startswith("Heading "):
            document.add_heading(text, level=int(style.split()[-1]))
        else:
            document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    document.save(str(path))
    return path


CONTAINER = (
    '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
    '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>'
)


def make_epub(
    chapters: list[tuple[str, str]],
    *,
    nav: bool = True,
    extra: dict[str, bytes] | None = None,
    title: str = "Test Book",
) -> bytes:
    """``chapters`` are ``(title, body_html)``. Builds an EPUB 3 package with an optional nav document."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
        book.writestr("META-INF/container.xml", CONTAINER)
        manifest = []
        spine = []
        nav_items = []
        for i, (chapter_title, body) in enumerate(chapters, start=1):
            name = f"chapter{i}.xhtml"
            book.writestr(
                f"OEBPS/{name}",
                f'<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>{chapter_title}</title>'
                f"<style>p{{color:red}}</style></head><body><h1>{chapter_title}</h1>{body}<script>alert(1)</script></body></html>",
            )
            manifest.append(f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="c{i}"/>')
            nav_items.append(f'<li><a href="{name}">{chapter_title}</a></li>')
        if nav:
            book.writestr(
                "OEBPS/nav.xhtml",
                '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body>'
                f'<nav epub:type="toc"><ol>{"".join(nav_items)}</ol></nav></body></html>',
            )
            manifest.append('<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>')
        book.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
            f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{title}</dc:title></metadata>'
            f"<manifest>{''.join(manifest)}</manifest><spine>{''.join(spine)}</spine></package>",
        )
        for name, data in (extra or {}).items():
            book.writestr(name, data)
    return buffer.getvalue()


def make_zip_bomb(kind: str = "docx") -> bytes:
    """A small archive that inflates to ~60 MB of zeros (ratio far above the limit)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        if kind == "epub":
            archive.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
            archive.writestr("META-INF/container.xml", CONTAINER)
        else:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", "<document/>")
        archive.writestr("bomb.bin", b"\0" * (60 * 1024 * 1024))
    return buffer.getvalue()


BILLION_LAUGHS = (
    '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;">'
    '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;">]><package>&lol3;</package>'
)
