# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Carson McRae - Local TTS Web UI - github.com/carsonmcrae/local-tts-web-ui
"""Purpose: Edge-case tests for the DOCX, EPUB, Markdown and HTML extractors, type sniffing and SafeZip.

Scope: DOCX Title and custom heading styles, empty paragraphs and table rows, damaged and empty files;
EPUB 2 NCX chapter titles, missing or non-HTML spine items, empty chapters and damaged package XML;
Markdown soft line breaks and empty headings; image alt text in chapter HTML; ZIP containers whose
contents do not match their extension; SafeZip size limits, missing and damaged members.
Does NOT belong: Happy paths (test_extractors.py) or hostile archives (test_document_security.py).
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE

from builders import CONTAINER, make_epub2, make_zip
from ltts.documents.extract_docx import extract_docx_file
from ltts.documents.extract_epub import extract_epub_file
from ltts.documents.extract_markdown import markdown_to_extraction
from ltts.documents.extractors import extract
from ltts.documents.html_text import html_to_blocks
from ltts.documents.safezip import SafeZip
from ltts.documents.sniff import detect_type
from ltts.documents.structures import DocumentType, Extraction, ExtractionLimits
from ltts.errors import DocumentExtractionError, UnsupportedDocumentError

LIMITS = ExtractionLimits(1_000_000, 50, 500, 20 * 1024 * 1024, 200)
XHTML = "application/xhtml+xml"


def write(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


def blocks_of(extraction: Extraction) -> list[tuple[str, list[tuple[str, str]]]]:
    return [(s.title, [(b.kind, b.text) for b in s.blocks]) for s in extraction.sections]


# -- DOCX -------------------------------------------------------------------------
def test_docx_title_style_custom_headings_and_empty_content(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("The Title Style", level=0)
    document.add_paragraph("Intro paragraph.")
    document.add_paragraph("")
    document.styles.add_style("Heading Custom", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Custom styled line.", style="Heading Custom")
    table = document.add_table(rows=2, cols=2)  # the first row stays empty
    table.cell(1, 0).text = "Kept"
    document.save(str(tmp_path / "a.docx"))
    result = extract_docx_file(tmp_path / "a.docx", "fallback", LIMITS)
    assert result.title == "fallback"
    assert blocks_of(result) == [
        (
            "The Title Style",
            [
                ("heading", "The Title Style"),
                ("paragraph", "Intro paragraph."),
                ("paragraph", "Custom styled line."),  # "Heading Custom" has no level number
                ("paragraph", "Kept"),
            ],
        )
    ]


def test_docx_without_text_has_no_sections(tmp_path: Path) -> None:
    docx.Document().save(str(tmp_path / "empty.docx"))
    assert extract_docx_file(tmp_path / "empty.docx", "empty", LIMITS).sections == []
    with pytest.raises(DocumentExtractionError) as error:
        extract(tmp_path / "empty.docx", DocumentType.DOCX, "empty", LIMITS)
    assert error.value.code == "empty_document"


def test_damaged_docx_is_reported(tmp_path: Path) -> None:
    path = write(tmp_path, "bad.docx", make_zip({"[Content_Types].xml": "<Types/>", "word/document.xml": "<w:doc"}))
    with pytest.raises(DocumentExtractionError) as error:
        extract_docx_file(path, "bad", LIMITS)
    assert error.value.code == "corrupt_document"


# -- EPUB -------------------------------------------------------------------------
def test_epub2_uses_ncx_titles_and_skips_unusable_spine_items(tmp_path: Path) -> None:
    ncx = (
        '<navPoint><navLabel><text>Chapter One</text></navLabel><content src="c1.xhtml"/></navPoint>'
        "<navPoint><navLabel><text>No source</text></navLabel></navPoint>"
        '<navPoint><navLabel><text></text></navLabel><content src="c2.xhtml"/></navPoint>'
    )
    data = make_epub2(
        [("c1", "c1.xhtml", XHTML), ("cover", "cover.png", "image/png"), ("c2", "c2.xhtml", XHTML)],
        ncx,
        manifest_extra=f'<item id="gone" href="missing.xhtml" media-type="{XHTML}"/>',
    )
    result = extract_epub_file(write(tmp_path, "old.epub", data), "fallback", LIMITS)
    assert result.title == "Old Book"
    # An NCX label names chapter one; chapter two has an empty label, so its first heading is used.
    assert [s.title for s in result.sections] == ["Chapter One", "Heading c2"]


def test_epub_nav_links_without_target_or_text_and_empty_chapters_are_skipped(tmp_path: Path) -> None:
    nav = (
        '<html xmlns="http://www.w3.org/1999/xhtml"><body><nav><ol>'
        '<li><a href="c1.xhtml">Named</a></li><li><a>No target</a></li><li><a href="c2.xhtml"> </a></li>'
        "</ol></nav></body></html>"
    )
    opf = (
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0"><metadata/><manifest>'
        f'<item id="c1" href="c1.xhtml" media-type="{XHTML}"/><item id="c2" href="c2.xhtml" media-type="{XHTML}"/>'
        f'<item id="nav" href="nav.xhtml" media-type="{XHTML}" properties="nav"/>'
        '</manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>'
    )
    data = make_zip(
        {
            "mimetype": "application/epub+zip",
            "META-INF/container.xml": CONTAINER,
            "OEBPS/content.opf": opf,
            "OEBPS/nav.xhtml": nav,
            "OEBPS/c1.xhtml": "<html><body><p>Text.</p></body></html>",
            "OEBPS/c2.xhtml": "<html><body><script>only(code)</script></body></html>",
        }
    )
    result = extract_epub_file(write(tmp_path, "n.epub", data), "fallback", LIMITS)
    assert result.title == "fallback"  # the package has no dc:title
    assert blocks_of(result) == [("Named", [("paragraph", "Text.")])]  # the empty second chapter is dropped


@pytest.mark.parametrize(
    "container",
    [
        "<container><rootfiles><rootfile full-path=",  # not well-formed XML
        CONTAINER.replace("OEBPS/content.opf", "OEBPS/missing.opf"),  # points at no package file
    ],
)
def test_epub_without_a_readable_package_is_reported(tmp_path: Path, container: str) -> None:
    path = write(
        tmp_path, "broken.epub", make_zip({"mimetype": "application/epub+zip", "META-INF/container.xml": container})
    )
    with pytest.raises(DocumentExtractionError) as error:
        extract_epub_file(path, "broken", LIMITS)
    assert error.value.code == "corrupt_document"


# -- Markdown and HTML -------------------------------------------------------------
def test_markdown_soft_breaks_join_and_empty_headings_are_skipped() -> None:
    result = markdown_to_extraction("Line one\nline two\n\n# ![](x.png)\n\nAfter.", "Given")
    assert result.title == "Given"  # the first section does not open with a heading
    assert blocks_of(result) == [("Given", [("paragraph", "Line one line two"), ("paragraph", "After.")])]


def test_meaningful_image_alt_text_is_read() -> None:
    markup = '<p>See <img alt="A lighthouse at dusk"> and <img alt="x"> <img></p><nav><img alt="Hidden alt text"></nav>'
    assert [b.text for b in html_to_blocks(markup)] == ["See A lighthouse at dusk and"]


# -- Sniffing and SafeZip -------------------------------------------------------------
@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("plain.epub", make_zip({"mimetype": "application/zip", "readme.txt": "hello"})),
        ("broken.docx", b"PK\x03\x04 this is not really a zip archive"),
    ],
)
def test_zip_containers_that_are_neither_epub_nor_docx_are_rejected(tmp_path: Path, name: str, data: bytes) -> None:
    with pytest.raises(UnsupportedDocumentError) as error:
        detect_type(write(tmp_path, name, data), name)
    assert error.value.code == "content_mismatch"


def test_safezip_rejects_non_zip_files(tmp_path: Path) -> None:
    with pytest.raises(DocumentExtractionError) as error:
        SafeZip(write(tmp_path, "x.zip", b"not a zip"), LIMITS)
    assert error.value.code == "corrupt_document"


def test_safezip_limits_the_total_uncompressed_size(tmp_path: Path) -> None:
    data = make_zip({"a.txt": "a" * 600, "b.txt": "b" * 600})
    with pytest.raises(DocumentExtractionError) as error:
        SafeZip(write(tmp_path, "big.zip", data), ExtractionLimits(1_000_000, 50, 500, 1_000, 200))
    assert error.value.code == "archive_limits"


def test_safezip_reports_missing_and_damaged_members(tmp_path: Path) -> None:
    payload = b"original member content"
    data = make_zip({"part.xml": payload}).replace(payload, b"tampered member content")
    with SafeZip(write(tmp_path, "d.zip", data), LIMITS) as archive:
        with pytest.raises(DocumentExtractionError, match="missing") as missing:
            archive.read("absent.xml")
        with pytest.raises(DocumentExtractionError, match="damaged") as damaged:
            archive.read("part.xml")  # the CRC no longer matches
    assert missing.value.code == damaged.value.code == "corrupt_document"
